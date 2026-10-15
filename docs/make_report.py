"""Build docs/report.docx from the real pipeline outputs in data/ (Gold tables, ALS metrics, run log).

    python docs/make_report.py
"""
import glob
import json
import re
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DATA, FIG = ROOT / "data", ROOT / "docs" / "figures"
FIG.mkdir(parents=True, exist_ok=True)
load = lambda rel: pd.read_parquet(DATA / rel)
metrics = json.loads(open(glob.glob(str(DATA / "output" / "metrics" / "part-*"))[0]).read())
kpis, stats, users = load("gold/yearly_kpis").sort_values("rating_year"), load("gold/movie_stats"), load("gold/user_activity")
gt, dist, dq = load("gold/genre_trends"), load("gold/rating_distribution"), load("silver/dq_report")
top_genre, recs, fav = load("gold/top_movies_by_genre"), load("output/user_recs_top10"), load("gold/user_top_rated")
pred = load("output/test_predictions_sample")
log = "\n".join(f.read_bytes().replace(b"\x00", b"").decode("utf-8", "ignore")  # full run first, ALS re-run last
                for f in sorted(DATA.glob("dataproc_run*.log"), reverse=True))
job_times = dict(re.findall(r"(\w+)\.py finished in ([\d,]+)s", log))


def fig(name, draw, w=7, h=3.4):
    plt.figure(figsize=(w, h))
    draw()
    plt.tight_layout()
    plt.savefig(FIG / f"{name}.png", dpi=150)
    plt.close()
    return FIG / f"{name}.png"


figs = {
    "yearly": fig("yearly", lambda: (plt.bar(kpis.rating_year, kpis.num_ratings / 1e6, color="#4C78A8"),
                                     plt.title("Ratings per year (millions)"), plt.xlabel("Year"))),
    "dist": fig("dist", lambda: (plt.bar(dist.rating.astype(str), dist["count"] / 1e6, color="#F58518"),
                                 plt.title("Rating distribution (millions)"), plt.xlabel("Rating"))),
    "genres": fig("genres", lambda: [plt.plot(d.rating_year, d.num_ratings / 1e3, label=g, marker=".")
                                     for g, d in gt[gt.genre.isin(gt.groupby("genre").num_ratings.sum().nlargest(6).index)]
                                     .sort_values("rating_year").groupby("genre")]
                  + [plt.legend(fontsize=7), plt.title("Ratings per year by genre (thousands, top 6 genres)")]),
    "grid": fig("grid", lambda: (plt.bar([f"rank={g['rank']}\nreg={g['regParam']}" for g in metrics["grid"]],
                                         [g["val_rmse"] for g in metrics["grid"]], color="#54A24B"),
                                 plt.ylim(min(g["val_rmse"] for g in metrics["grid"]) - 0.05,
                                          max(g["val_rmse"] for g in metrics["grid"]) + 0.02),
                                 plt.title("ALS hyper-parameter search (validation RMSE, lower is better)"))),
    "pred": fig("pred", lambda: (pred.boxplot(column="prediction", by="rating", grid=False), plt.suptitle(""),
                                 plt.title("Predicted vs actual rating (test sample)"), plt.xlabel("Actual rating"),
                                 plt.ylabel("Predicted"))),
}

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def fmt(v):
    if isinstance(v, (int, np.integer)):
        return f"{v:,}"
    if isinstance(v, (float, np.floating)):
        if v != v:
            return "–"
        return f"{v:.0f}" if v.is_integer() and abs(v) > 100 else f"{v:.3f}"  # years as 1994, ratings as 4.500
    return str(v)


def table(df, cols=None):
    df = df[cols] if cols else df
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for row in df.itertuples(index=False):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = fmt(v)
    doc.add_paragraph()


def image(key, caption):
    doc.add_picture(str(figs[key]), width=Inches(6.2))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True


bullets = lambda items: [doc.add_paragraph(i, style="List Bullet") for i in items]
n_ratings, n_users, n_movies = int(kpis.num_ratings.sum()), len(users), len(stats)

title = doc.add_heading("Movie Recommendation System using Big Data Analytics", 0)
doc.add_paragraph(f"Minor Project Report  ·  Krishanu Anand  ·  {date.today():%d %B %Y}")
doc.add_paragraph("GitHub: https://github.com/anand-krishanu/AttractorField")

doc.add_heading("Abstract", 1)
doc.add_paragraph(
    f"This project builds an end-to-end Big Data pipeline on the MovieLens {metrics['dataset']} dataset "
    f"({n_ratings:,} ratings from {n_users:,} users on {n_movies:,} movies). Data is ingested into Google Cloud "
    "Storage and processed with PySpark on Google Cloud Dataproc using a Bronze → Silver → Gold medallion architecture. "
    "Spark SQL and window functions produce analysis-ready tables, and a Spark MLlib ALS collaborative-filtering model "
    f"predicts ratings with a test RMSE of {metrics['test_rmse']} (vs. {metrics['baseline_rmse_global_mean']} for a "
    "global-mean baseline) and generates a personalised Top-10 list for every user. Results are presented in an "
    "interactive Streamlit dashboard, and the workflow is orchestrated with a runner script and an Airflow DAG.")

doc.add_heading("1. Problem Definition", 1)
doc.add_paragraph(
    "Streaming catalogues contain tens of thousands of titles; users struggle to discover movies they will enjoy, and "
    "content teams lack visibility into how tastes evolve. Manually curated lists do not scale to millions of users.")
doc.add_heading("Objectives", 2)
bullets(["Ingest and store a large real-world ratings dataset in the cloud with full lineage metadata.",
         "Clean and validate the data (nulls, duplicates, invalid values, orphan keys, inconsistent text).",
         "Produce analysis-ready Gold tables: popularity, weighted ratings, genre trends, user segments, KPIs.",
         "Train and evaluate a scalable recommender (ALS) that beats a naive baseline.",
         "Deliver personalised Top-10 recommendations and an interactive dashboard built on real outputs."])
doc.add_heading("Target users & expected outcomes", 2)
bullets(["Viewers: personalised, unseen movie recommendations and 'similar movies'.",
         "Product/content teams: KPIs, genre trends and user segments to guide acquisition and promotion.",
         "Outcome: an automated, reproducible pipeline from raw data source to end-user dashboard."])

doc.add_heading("2. Dataset", 1)
doc.add_paragraph(
    "MovieLens 25M is a stable benchmark dataset released by GroupLens Research (University of Minnesota) in 2019. "
    "It contains explicit 0.5–5 star ratings collected on movielens.org between 1995 and 2019; each user has rated "
    "at least 20 movies. Files: ratings.csv (userId, movieId, rating, timestamp), movies.csv (title with year, "
    "pipe-separated genres), tags.csv (free-text user tags) and links.csv (IMDb/TMDb ids). "
    "Source: https://grouplens.org/datasets/movielens/25m/ . The smaller ml-latest-small release (100k ratings) was "
    "used for local development and testing.")
doc.add_paragraph("Citation: F. Maxwell Harper and Joseph A. Konstan. 2015. The MovieLens Datasets: History and "
                  "Context. ACM Transactions on Interactive Intelligent Systems 5, 4, Article 19. "
                  "https://doi.org/10.1145/2827872")

doc.add_heading("3. Architecture & Technologies", 1)
doc.add_paragraph(
    "GroupLens → ingest.py → GCS raw/ → Bronze (Parquet) → Silver (clean, partitioned) → Gold (analytics) + "
    "MLlib ALS (output/) → Streamlit dashboard. Orchestration: run_pipeline.py (local Docker), "
    "scripts/dataproc_run.ps1 (Google Cloud) and an Airflow DAG (dags/movie_reco_dag.py).")
table(pd.DataFrame([
    ["Storage", "Google Cloud Storage (data lake), Parquet (columnar, partitioned)"],
    ["Processing", "Apache Spark 3.5 / PySpark, Spark SQL, window functions"],
    ["Machine learning", "Spark MLlib ALS, RegressionEvaluator, RankingEvaluator"],
    ["Cloud compute", "Google Cloud Dataproc 2.2 (1 × e2-standard-2 master, 2 × e2-highmem-4 workers)"],
    ["Orchestration", "Apache Airflow DAG (Dataproc operators), PowerShell/Python runners"],
    ["Visualisation", "Streamlit + Plotly dashboard"],
    ["Dev environment", "Docker image apache/spark 3.5.3 (Java 17) mirroring Dataproc; GitHub with PR workflow"],
], columns=["Layer", "Technology"]))

doc.add_heading("4. Methodology", 1)
doc.add_heading("4.1 Ingestion & Bronze layer", 2)
doc.add_paragraph(
    "ingest.py downloads the official zip, writes a manifest (source URL, timestamp, file sizes) and uploads the raw "
    "CSVs to GCS — on Dataproc this runs on the cluster so data never passes through the client. bronze.py reads the "
    "CSVs with explicit schemas (no inference) and adds ingest_ts, source_file and dataset_version columns before "
    "writing Parquet, preserving the raw values for lineage and reprocessing.")
doc.add_heading("4.2 Silver layer – cleaning", 2)
doc.add_paragraph(
    "silver.py removes null keys, out-of-range ratings and invalid timestamps; resolves duplicate (user, movie) "
    "ratings by keeping the latest one with a row_number() window; drops ratings that reference unknown movies "
    "(left-semi join); converts epoch seconds to timestamps; extracts the release year from the title with a regex; "
    "turns the pipe-separated genres into arrays (mapping '(no genres listed)' to null); lower-cases and de-duplicates "
    "tags; and joins IMDb/TMDb links. Ratings are written partitioned by rating_year. Each check is logged:")
table(dq, ["table", "check", "rows_before", "rows_after", "removed"])
doc.add_heading("4.3 Gold layer – analytics", 2)
bullets(["movie_stats: count, mean, std-dev and an IMDb-style Bayesian weighted rating "
         "WR = v/(v+m)·R + m/(v+m)·C (m = 90th percentile of vote counts) so that low-count movies do not dominate.",
         "top_movies_by_genre: rank() over a window partitioned by genre (top 20 per genre).",
         "genre_trends: Spark SQL over a temp view with LAG() to compute year-over-year growth per genre.",
         "user_activity: per-user volume, average rating, first/last activity and a segment (casual→power).",
         "user_top_rated, rating_distribution, top_tags and yearly_kpis (with a running cumulative total)."])
doc.add_heading("4.4 Recommendation model – Spark MLlib ALS", 2)
doc.add_paragraph(
    "Alternating Least Squares factorises the sparse user × movie rating matrix into low-rank user and item factor "
    "matrices; predicted rating = dot product of the two factor vectors. A first run on all movies showed that "
    "titles with only 2–10 ratings received extreme factors and flooded every user's top-10 (precision@10 = 0), "
    f"so the model is trained only on movies with at least {metrics.get('min_movie_ratings', 50)} ratings "
    f"({metrics.get('movies_in_model', 0):,} movies, still the vast majority of all ratings). Data was split 70/10/20 (train/validation/"
    "test, seed 42). A grid over rank ∈ {10, 20} and regParam ∈ {0.05, 0.1} (maxIter 10, non-negative factors, "
    "coldStartStrategy=drop) was evaluated on the validation split; the best configuration was refit on train+val "
    "and evaluated once on the held-out test split. Ranking quality was measured with precision@10 and NDCG@10 on a "
    "sample of 2,000 users, treating test movies rated ≥ 4 as relevant. Recommendations exclude movies the user has "
    "already rated. Item-item similarity was computed as cosine similarity between ALS item factors using Spark SQL "
    "higher-order functions.")

doc.add_heading("5. Results & Analysis", 1)
doc.add_heading("5.1 Exploratory analytics", 2)
image("yearly", "Figure 1 – Rating volume per year")
image("dist", "Figure 2 – Distribution of ratings")
image("genres", "Figure 3 – Genre popularity over time")
r = dist.set_index("rating")["count"]
doc.add_paragraph(
    f"Ratings are skewed positive: {r[r.index >= 4].sum() / r.sum():.0%} of ratings are 4 stars or more and whole-star "
    f"ratings dominate half-stars. The busiest year was {int(kpis.loc[kpis.num_ratings.idxmax(), 'rating_year'])}. "
    f"Users rated on average {users.num_ratings.mean():.0f} movies (median {users.num_ratings.median():.0f}), "
    f"showing a long-tail activity distribution; {(users.segment == 'power').mean():.1%} of users are 'power' users "
    "with 500+ ratings.")
doc.add_heading("Top-rated movies (Drama, by weighted rating)", 3)
table(top_genre[top_genre.genre == "Drama"].sort_values("rank").head(10),
      ["rank", "clean_title", "release_year", "num_ratings", "avg_rating", "weighted_rating"])

doc.add_heading("5.2 Model evaluation", 2)
table(pd.DataFrame([[g["rank"], g["regParam"], g["val_rmse"]] for g in metrics["grid"]],
                   columns=["rank", "regParam", "validation RMSE"]))
image("grid", "Figure 4 – Hyper-parameter search")
table(pd.DataFrame([
    ["Best parameters", f"rank={metrics['best_rank']}, regParam={metrics['best_regParam']}"],
    ["Train / Val / Test rows", f"{metrics['train_rows']:,} / {metrics['val_rows']:,} / {metrics['test_rows']:,}"],
    ["Test RMSE (ALS)", metrics["test_rmse"]],
    ["Test RMSE (global-mean baseline)", metrics["baseline_rmse_global_mean"]],
    ["Test MAE", metrics["test_mae"]],
    ["Precision@10", metrics["precision_at_10"]],
    ["NDCG@10", metrics["ndcg_at_10"]],
], columns=["Metric", "Value"]))
gain = 1 - metrics["test_rmse"] / metrics["baseline_rmse_global_mean"]
doc.add_paragraph(
    f"ALS reduces RMSE by {gain:.1%} relative to always predicting the global mean, i.e. on average its predictions "
    f"are within ~{metrics['test_mae']:.2f} stars of the true rating. Ranking metrics are low in absolute terms, "
    "which is expected for top-10 retrieval from tens of thousands of movies when only held-out test ratings count as "
    "relevant; they serve as a baseline for future improvements.")
image("pred", "Figure 5 – Predicted vs actual ratings on the test split")

doc.add_heading("5.3 Sample recommendations", 2)
uid = int(users[(users.num_ratings > 100) & users.userId.isin(recs.userId)].userId.iloc[0])
doc.add_paragraph(f"User {uid} – favourite movies (highest rated):")
table(fav[fav.userId == uid].sort_values("rank").head(5), ["clean_title", "rating"])
doc.add_paragraph(f"User {uid} – Top-10 recommendations (unseen):")
table(recs[recs.userId == uid].sort_values("rank"), ["rank", "clean_title", "release_year", "predicted_rating"])

doc.add_heading("6. Dashboard", 1)
doc.add_paragraph(
    "The Streamlit dashboard (dashboard/app.py) reads the Gold and model output Parquet files produced by the "
    "Dataproc run. It was designed from business requirements: each page answers the question of one stakeholder "
    "group, and every page follows the same reading order of headline numbers, then the main chart, then detail. "
    "Chart titles state the finding rather than the metric, one accent colour marks what matters, numbers are "
    "abbreviated for readability, and labels use plain language. Run: streamlit run dashboard/app.py")
table(pd.DataFrame([
    ["Overview", "Leadership", "Platform health and recommender impact"],
    ["Content", "Content and acquisition", "Top movies per genre, genre volume vs rating, trends, tags, similar movies"],
    ["Audience", "Growth and marketing", "Engagement segments, active users over time, rating behaviour"],
    ["Recommendations", "Personalisation", "A user's favourites next to their predicted Top-10"],
    ["Model quality", "Data science", "RMSE vs baseline, tuning results, error analysis"],
    ["Data pipeline", "Data engineering", "Medallion layers, data-quality checks, Dataproc job runtimes"],
], columns=["Page", "Audience", "What it answers"]))
shots = sorted((ROOT / "docs" / "dashboard").glob("*.png"))
for i, shot in enumerate(shots, 1):
    doc.add_picture(str(shot), width=Inches(6.2))
    p = doc.add_paragraph(f"Dashboard {i} – {shot.stem.split('_', 1)[1].replace('_', ' ').capitalize()} page")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True

doc.add_heading("7. Cloud execution & orchestration", 1)
doc.add_paragraph(
    "scripts/dataproc_run.ps1 creates a GCS bucket and a Dataproc 2.2 cluster (auto-deleted after 30 minutes idle), "
    "uploads the code, submits ingest, bronze, silver, gold and train_als as PySpark jobs, downloads the outputs for "
    "the dashboard and deletes the cluster to minimise cost. The same steps are defined as an Airflow DAG "
    "(create_cluster → ingest → bronze → silver → gold → train_als → delete_cluster, weekly schedule, retries, "
    "cluster deletion with trigger_rule=ALL_DONE) for deployment on Cloud Composer.")
if job_times and not (ROOT / "docs" / "cloud_run_evidence" / "jobs.txt").exists():
    table(pd.DataFrame(list(job_times.items()), columns=["Dataproc job", "Duration (s)"]))
evidence = ROOT / "docs" / "cloud_run_evidence"
if (evidence / "jobs.txt").exists():
    doc.add_heading("Cloud execution evidence", 2)
    doc.add_paragraph(
        "The repository folder docs/cloud_run_evidence/ contains the Dataproc job history (jobs.txt), the cleaned run "
        "logs, the GCS bucket listing per layer (bucket_listing.txt), the Spark driver output of the final ALS job and "
        "GCP console screenshots. Job history (project movie-reco-bigdata-2026, region asia-south1):")
    jl = [l.split() for l in (evidence / "jobs.txt").read_text(encoding="utf-8").splitlines() if re.match(r"^[0-9a-f]{32}\s", l)]
    table(pd.DataFrame([[r[0][:8], r[1], r[-5], r[-4] + " " + r[-3], r[-1]] for r in jl],
                       columns=["Job id", "Script", "State", "Start (UTC)", "Seconds"]))
    sizes = re.findall(r"^([\d.]+\s*\w+)\s+gs://\S+/(\w+)$", (evidence / "bucket_listing.txt").read_text(encoding="utf-8"), re.M)
    if sizes:
        table(pd.DataFrame([(layer, size) for size, layer in sizes], columns=["GCS layer", "Size"]))
    for shot in sorted((evidence / "screenshots").glob("*.png")):
        doc.add_picture(str(shot), width=Inches(6.2))
        doc.add_paragraph(shot.stem.replace("_", " ")).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    "Development followed a Git feature-branch workflow: each part (setup/ingest, Bronze/Silver, Gold/ML, cloud "
    "orchestration, dashboard/docs) was developed on its own branch and merged through a pull request.")

doc.add_heading("8. Conclusion & Future Work", 1)
doc.add_paragraph(
    f"The project demonstrates scalable Big Data processing of {n_ratings:,} ratings on Google Cloud: a governed "
    "medallion data lake, Spark SQL analytics, a tuned MLlib recommender that clearly beats the baseline, and an "
    "interactive dashboard on real outputs. Future work: implicit-feedback ALS and hybrid content features "
    "(genres, tags) for cold-start movies, larger hyper-parameter search with CrossValidator, Spark Structured "
    "Streaming for real-time rating ingestion, and serving recommendations through an API.")

doc.add_heading("References", 1)
bullets(["F. M. Harper and J. A. Konstan. The MovieLens Datasets: History and Context. ACM TiiS 5(4), 2015.",
         "Y. Koren, R. Bell, C. Volinsky. Matrix Factorization Techniques for Recommender Systems. IEEE Computer, 2009.",
         "Y. Hu, Y. Koren, C. Volinsky. Collaborative Filtering for Implicit Feedback Datasets. ICDM, 2008.",
         "Apache Spark 3.5 MLlib – Collaborative Filtering: https://spark.apache.org/docs/3.5.3/ml-collaborative-filtering.html",
         "Google Cloud Dataproc documentation: https://cloud.google.com/dataproc/docs",
         "GroupLens MovieLens 25M: https://grouplens.org/datasets/movielens/25m/"])

out = ROOT / "docs" / "report.docx"
doc.save(out)
print(f"Wrote {out}")
